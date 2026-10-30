import datetime
import json
import time
from io import BytesIO

from openpyxl import load_workbook
from unittest.mock import patch
from django.db import IntegrityError, connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from accounts.models import User
from notifications.models import AuditLog
from students.models import AcademicClass, Section, Student

from . import services
from .models import Attendance

TODAY = "2026-10-23"


class AttendanceTestBase(TestCase):
    """Shared setup: one class with 3 students, its teacher, and another teacher."""

    def setUp(self):
        self.teacher = User.objects.create_user(
            username="teacher1", password="pass12345", role=User.Role.TEACHER
        )
        self.other_teacher = User.objects.create_user(
            username="teacher2", password="pass12345", role=User.Role.TEACHER
        )
        self.admin = User.objects.create_user(
            username="admin1", password="pass12345", role=User.Role.ADMIN
        )
        self.cls = AcademicClass.objects.create(name="Grade 10", class_teacher=self.teacher)
        section = Section.objects.create(academic_class=self.cls, name="A")
        self.students = [
            Student.objects.create(
                name=f"Student {n}",
                student_class="Grade 10",
                roll_number=n,
                contact_details=f"parent{n}@example.com",
                section=section,
            )
            for n in (1, 2, 3)
        ]

    def post_json(self, url, payload):
        return self.client.post(url, json.dumps(payload), content_type="application/json")

    def mark_payload(self, records=None, date=TODAY):
        if records is None:
            records = [{"studentId": s.pk, "status": "Present"} for s in self.students]
        return {"classId": self.cls.pk, "date": date, "records": records}


# Unit tests

class AttendanceServiceUnitTests(AttendanceTestBase):
    def test_valid_statuses_are_accepted(self):
        for status in ("Present", "Absent", "Late", "Excused"):
            self.assertTrue(services.is_valid_status(status))

    def test_invalid_status_is_rejected(self):
        self.assertFalse(services.is_valid_status("Sick"))
        self.assertFalse(services.is_valid_status(""))

    def test_duplicate_check_false_before_marking(self):
        self.assertFalse(services.attendance_exists(self.cls, datetime.date(2026, 10, 23)))

    def test_duplicate_check_true_after_marking(self):
        services.mark_attendance(self.teacher, self.cls.pk, TODAY, [])
        self.assertTrue(services.attendance_exists(self.cls, datetime.date(2026, 10, 23)))

    def test_only_assigned_teacher_counts_as_class_teacher(self):
        self.assertTrue(services.is_class_teacher(self.teacher, self.cls))
        self.assertFalse(services.is_class_teacher(self.other_teacher, self.cls))
        self.assertFalse(services.is_class_teacher(self.admin, self.cls))

    def test_timestamp_set_on_save_and_updated_on_edit(self):
        entries = services.mark_attendance(self.teacher, self.cls.pk, TODAY, [])
        record = Attendance.objects.get(pk=entries[0].pk)
        first_update = record.updated_at
        self.assertIsNotNone(record.created_at)
        services.edit_attendance(self.teacher, record, "Late")
        record.refresh_from_db()
        self.assertGreater(record.updated_at, first_update)

    def test_database_blocks_duplicate_rows(self):
        Attendance.objects.create(
            student=self.students[0], academic_class=self.cls,
            date=datetime.date(2026, 10, 23), status="Present", marked_by=self.teacher,
        )
        with self.assertRaises(IntegrityError):
            Attendance.objects.create(
                student=self.students[0], academic_class=self.cls,
                date=datetime.date(2026, 10, 23), status="Absent", marked_by=self.teacher,
            )


# Integration tests:

class MarkAttendanceApiTests(AttendanceTestBase):
    url = "/api/attendance"

    def test_teacher_marks_whole_class(self):  
        self.client.force_login(self.teacher)
        response = self.post_json(self.url, self.mark_payload())
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["savedCount"], 3)
        self.assertEqual(Attendance.objects.count(), 3)

    def test_all_four_statuses_are_saved(self):
        self.client.force_login(self.teacher)
        records = [
            {"studentId": self.students[0].pk, "status": "Absent"},
            {"studentId": self.students[1].pk, "status": "Late"},
            {"studentId": self.students[2].pk, "status": "Excused", "remarks": "Doctor"},
        ]
        self.post_json(self.url, self.mark_payload(records))
        saved = {a.student_id: a.status for a in Attendance.objects.all()}
        self.assertEqual(saved[self.students[0].pk], "Absent")
        self.assertEqual(saved[self.students[1].pk], "Late")
        self.assertEqual(saved[self.students[2].pk], "Excused")
        self.assertEqual(Attendance.objects.get(student=self.students[2]).remarks, "Doctor")

    def test_mark_all_present_with_exceptions(self):  
        self.client.force_login(self.teacher)
        records = [{"studentId": self.students[1].pk, "status": "Absent"}]
        response = self.post_json(self.url, self.mark_payload(records))
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Attendance.objects.filter(status="Present").count(), 2)
        self.assertEqual(Attendance.objects.get(student=self.students[1]).status, "Absent")

    def test_marking_twice_returns_409(self):  
        self.client.force_login(self.teacher)
        self.post_json(self.url, self.mark_payload())
        response = self.post_json(self.url, self.mark_payload())
        self.assertEqual(response.status_code, 409)
        self.assertEqual(Attendance.objects.count(), 3)

    def test_same_class_on_different_date_is_allowed(self):
        self.client.force_login(self.teacher)
        self.post_json(self.url, self.mark_payload(date="2026-10-22"))
        response = self.post_json(self.url, self.mark_payload(date="2026-10-23"))
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Attendance.objects.count(), 6)

    def test_wrong_teacher_gets_403(self):
        self.client.force_login(self.other_teacher)
        response = self.post_json(self.url, self.mark_payload())
        self.assertEqual(response.status_code, 403)
        self.assertEqual(Attendance.objects.count(), 0)

    def test_admin_cannot_mark_attendance(self):
        self.client.force_login(self.admin)
        response = self.post_json(self.url, self.mark_payload())
        self.assertEqual(response.status_code, 403)

    def test_invalid_status_gets_400(self):
        self.client.force_login(self.teacher)
        records = [{"studentId": self.students[0].pk, "status": "Sick"}]
        response = self.post_json(self.url, self.mark_payload(records))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Attendance.objects.count(), 0)

    def test_student_from_another_class_gets_400(self):
        other_cls = AcademicClass.objects.create(name="Grade 9", class_teacher=self.teacher)
        other_section = Section.objects.create(academic_class=other_cls, name="B")
        outsider = Student.objects.create(
            name="Outsider", student_class="Grade 9", roll_number=99,
            contact_details="x@example.com", section=other_section,
        )
        self.client.force_login(self.teacher)
        records = [{"studentId": outsider.pk, "status": "Present"}]
        response = self.post_json(self.url, self.mark_payload(records))
        self.assertEqual(response.status_code, 400)

    def test_bad_date_gets_400(self):
        self.client.force_login(self.teacher)
        response = self.post_json(self.url, self.mark_payload(date="23-10-2026"))
        self.assertEqual(response.status_code, 400)

    def test_missing_fields_get_400(self):
        self.client.force_login(self.teacher)
        response = self.post_json(self.url, {"records": []})
        self.assertEqual(response.status_code, 400)

    def test_class_with_no_students_gets_400(self):
        empty = AcademicClass.objects.create(name="Empty", class_teacher=self.teacher)
        self.client.force_login(self.teacher)
        response = self.post_json(self.url, {"classId": empty.pk, "date": TODAY, "records": []})
        self.assertEqual(response.status_code, 400)

    def test_not_logged_in_gets_401(self):
        response = self.post_json(self.url, self.mark_payload())
        self.assertEqual(response.status_code, 401)


class EditAttendanceApiTests(AttendanceTestBase):
    def setUp(self):
        super().setUp()
        services.mark_attendance(self.teacher, self.cls.pk, TODAY, [])
        self.record = Attendance.objects.get(student=self.students[0])

    def edit_url(self):
        return reverse("attendance:edit", args=[self.record.pk])

    def test_teacher_edits_status_and_audit_log_is_created(self):  
        self.client.force_login(self.teacher)
        response = self.post_json(self.edit_url(), {"status": "Absent"})
        self.assertEqual(response.status_code, 200)
        self.record.refresh_from_db()
        self.assertEqual(self.record.status, "Absent")
        log = AuditLog.objects.get()
        self.assertEqual(log.edited_by, self.teacher)
        self.assertEqual((log.old_status, log.new_status), ("Present", "Absent"))

    def test_wrong_teacher_cannot_edit(self):
        self.client.force_login(self.other_teacher)
        response = self.post_json(self.edit_url(), {"status": "Absent"})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(AuditLog.objects.count(), 0)

    def test_invalid_status_on_edit_gets_400(self):
        self.client.force_login(self.teacher)
        response = self.post_json(self.edit_url(), {"status": "Gone"})
        self.assertEqual(response.status_code, 400)

    def test_editing_missing_record_gets_404(self):
        self.client.force_login(self.teacher)
        response = self.post_json(reverse("attendance:edit", args=[9999]), {"status": "Absent"})
        self.assertEqual(response.status_code, 404)

class DailyAttendanceApiTests(AttendanceTestBase):
    """REQ4.1: daily attendance view for a class and date."""

    url = "/api/attendance/daily"

    def mark_with_exceptions(self, date=TODAY):
        records = [
            {"studentId": self.students[1].pk, "status": "Absent"},
            {"studentId": self.students[2].pk, "status": "Late", "remarks": "Bus"},
        ]
        services.mark_attendance(self.teacher, self.cls.pk, date, records)

    def get_daily(self, class_id=None, date=TODAY):
        params = {"classId": self.cls.pk if class_id is None else class_id, "date": date}
        return self.client.get(self.url, params)

    def test_teacher_sees_every_student_status_for_a_marked_day(self):
        self.mark_with_exceptions()
        self.client.force_login(self.teacher)
        response = self.get_daily()
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertTrue(body["marked"])
        self.assertEqual(body["className"], "Grade 10")
        self.assertEqual(
            [(r["rollNumber"], r["status"]) for r in body["records"]],
            [(1, "Present"), (2, "Absent"), (3, "Late")],
        )
        self.assertEqual(body["records"][2]["remarks"], "Bus")
        self.assertEqual(
            body["counts"], {"Present": 1, "Absent": 1, "Late": 1, "Excused": 0}
        )

    def test_day_with_no_attendance_returns_clear_message(self):
        self.client.force_login(self.teacher)
        response = self.get_daily()
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertFalse(body["marked"])
        self.assertEqual(body["records"], [])
        self.assertIn("not been marked", body["message"])

    def test_only_the_requested_date_is_returned(self):
        self.mark_with_exceptions(date="2026-10-22")
        self.client.force_login(self.teacher)
        body = self.get_daily(date=TODAY).json()
        self.assertFalse(body["marked"])
        self.assertTrue(self.get_daily(date="2026-10-22").json()["marked"])

    def test_student_filter_returns_only_that_students_record_and_counts(self):
        self.mark_with_exceptions()
        self.client.force_login(self.teacher)
        response = self.client.get(
            self.url,
            {"classId": self.cls.pk, "date": TODAY, "studentId": self.students[1].pk},
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(len(body["records"]), 1)
        self.assertEqual(body["records"][0]["studentId"], self.students[1].pk)
        self.assertEqual(body["counts"]["Absent"], 1)

    def test_admin_can_view_any_class(self):
        self.mark_with_exceptions()
        self.client.force_login(self.admin)
        response = self.get_daily()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["records"]), 3)

    def test_other_teacher_gets_403(self):
        self.mark_with_exceptions()
        self.client.force_login(self.other_teacher)
        self.assertEqual(self.get_daily().status_code, 403)

    def test_student_parent_gets_403(self):
        parent = User.objects.create_user(
            username="parent1", password="pass12345", role=User.Role.STUDENT_PARENT
        )
        self.client.force_login(parent)
        self.assertEqual(self.get_daily().status_code, 403)

    def test_not_logged_in_is_refused(self):
        self.assertIn(self.get_daily().status_code, (401, 403))

    def test_bad_input_gets_400_or_404(self):
        self.client.force_login(self.teacher)
        self.assertEqual(self.get_daily(date="23-10-2026").status_code, 400)
        self.assertEqual(self.client.get(self.url, {"date": TODAY}).status_code, 400)
        self.assertEqual(self.client.get(self.url, {"classId": self.cls.pk}).status_code, 400)
        self.assertEqual(self.get_daily(class_id="abc").status_code, 400)
        self.assertEqual(self.get_daily(class_id=9999).status_code, 404)

class FiftyStudentSubmissionTests(AttendanceTestBase):
    """REQ3.5: a full class of 50 students is saved in one request, in under 5 seconds."""

    url = "/api/attendance"
    time_limit_seconds = 5

    def make_class(self, name, size, first_roll):
        academic_class = AcademicClass.objects.create(name=name, class_teacher=self.teacher)
        section = Section.objects.create(academic_class=academic_class, name="A")
        Student.objects.bulk_create(
            [
                Student(
                    name=f"{name} student {n}",
                    student_class=name,
                    roll_number=first_roll + n,
                    contact_details=f"parent{first_roll + n}@example.com",
                    section=section,
                )
                for n in range(size)
            ]
        )
        return academic_class

    def submit(self, academic_class, records):
        payload = {"classId": academic_class.pk, "date": TODAY, "records": records}
        return self.client.post(self.url, json.dumps(payload), content_type="application/json")

    def test_class_of_50_is_saved_in_one_request_under_5_seconds(self):
        academic_class = self.make_class("Grade 11", 50, first_roll=1000)
        students = list(Student.objects.filter(section__academic_class=academic_class))
        # A realistic mix: every 5th student is Absent, every 7th is Late
        records = []
        for index, student in enumerate(students, start=1):
            status = "Absent" if index % 5 == 0 else "Late" if index % 7 == 0 else "Present"
            records.append({"studentId": student.pk, "status": status})
        self.client.force_login(self.teacher)

        started = time.perf_counter()
        response = self.submit(academic_class, records)
        elapsed = time.perf_counter() - started

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["savedCount"], 50)
        self.assertEqual(Attendance.objects.filter(academic_class=academic_class).count(), 50)
        self.assertLess(elapsed, self.time_limit_seconds, f"Took {elapsed:.2f}s")

    def test_number_of_queries_does_not_grow_with_class_size(self):
        small = self.make_class("Small", 5, first_roll=100)
        large = self.make_class("Large", 50, first_roll=200)
        self.client.force_login(self.teacher)

        with CaptureQueriesContext(connection) as small_queries:
            self.assertEqual(self.submit(small, []).status_code, 201)
        with CaptureQueriesContext(connection) as large_queries:
            self.assertEqual(self.submit(large, []).status_code, 201)

        # 50 students must cost no more queries than 5 (no one-query-per-student).
        self.assertEqual(len(small_queries), len(large_queries))

class ReliabilityTests(AttendanceTestBase):
    """NFR3: no conflicting entries, and the system can be checked for uptime."""

    url = "/api/attendance"

    def test_two_submissions_at_the_same_moment_save_only_one_set(self):
        self.client.force_login(self.teacher)
        first_records = [{"studentId": self.students[0].pk, "status": "Absent"}]

        # Simulate the race: both requests pass the "already marked?" check,
        # so only the database unique constraint can stop the second one.
        with patch("attendance.services.attendance_exists", return_value=False):
            first = self.post_json(self.url, self.mark_payload(first_records))
            second = self.post_json(self.url, self.mark_payload())

        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 409)
        self.assertEqual(
            Attendance.objects.filter(academic_class=self.cls, date=TODAY).count(),
            len(self.students),
        )
        # The first submission is untouched by the rejected second one
        self.assertEqual(Attendance.objects.get(student=self.students[0]).status, "Absent")

    def test_rejected_duplicate_leaves_no_partial_data(self):
        self.client.force_login(self.teacher)
        self.post_json(self.url, self.mark_payload())
        before = list(Attendance.objects.order_by("pk").values_list("pk", "status"))
        self.post_json(self.url, self.mark_payload([{"studentId": self.students[1].pk, "status": "Late"}]))
        after = list(Attendance.objects.order_by("pk").values_list("pk", "status"))
        self.assertEqual(before, after)

    def test_health_check_reports_ok_without_login(self):
        response = self.client.get("/health/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "ok", "database": "up"})

    def test_health_check_reports_503_when_database_is_down(self):
        with patch("attendance.views.connection.cursor", side_effect=Exception("db down")):
            response = self.client.get("/health/")
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["status"], "error")

class ScaleAndIndexTests(AttendanceTestBase):
    """NFR1: the Attendance table is indexed and reports stay fast with 1000+ students."""

    def index_column_sets(self):
        with connection.cursor() as cursor:
            constraints = connection.introspection.get_constraints(
                cursor, Attendance._meta.db_table
            )
        return [
            tuple(info["columns"])
            for info in constraints.values()
            if info["index"] or info["unique"]
        ]

    def test_attendance_table_is_indexed_on_student_class_and_date(self):
        indexed = self.index_column_sets()
        # One index covering student + class + date (the REQ3.4 unique rule) ...
        self.assertIn(("student_id", "academic_class_id", "date"), indexed)
        # ... and one for "who was marked in this class on this date" reports
        self.assertIn(("academic_class_id", "date"), indexed)

    def test_daily_view_stays_fast_with_1000_students_of_data(self):
        # 20 classes x 50 students = 1000 students, each with 3 days of attendance
        admin_classes = []
        roll = 5000
        for number in range(20):
            academic_class = AcademicClass.objects.create(
                name=f"Scale {number}", class_teacher=self.teacher
            )
            section = Section.objects.create(academic_class=academic_class, name="A")
            Student.objects.bulk_create(
                [
                    Student(
                        name=f"Scale student {number}-{n}",
                        student_class=academic_class.name,
                        roll_number=roll + n,
                        contact_details=f"scale{number}-{n}@example.com",
                        section=section,
                    )
                    for n in range(50)
                ]
            )
            roll += 100
            admin_classes.append(academic_class)

        rows = [
            Attendance(student=student, academic_class=academic_class, date=day, status="Present")
            for academic_class in admin_classes
            for student in Student.objects.filter(section__academic_class=academic_class)
            for day in ("2026-10-20", "2026-10-21", "2026-10-22")
        ]
        Attendance.objects.bulk_create(rows)
        self.assertEqual(Student.objects.count(), 1000 + len(self.students))
        self.assertEqual(Attendance.objects.count(), 3000)

        self.client.force_login(self.teacher)
        started = time.perf_counter()
        response = self.client.get(
            "/api/attendance/daily",
            {"classId": admin_classes[7].pk, "date": "2026-10-21"},
        )
        elapsed = time.perf_counter() - started

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["records"]), 50)
        self.assertLess(elapsed, 2, f"Daily view took {elapsed:.2f}s")

class AttendanceSummaryUnitTests(TestCase):
    """REQ4.2: percentage = days present / total school days."""

    def test_no_days_gives_zero_and_does_not_crash(self):
        self.assertEqual(
            services.calculate_attendance_summary([]),
            {"totalDays": 0, "presentDays": 0, "percentage": 0.0},
        )

    def test_all_present_is_100_percent(self):
        summary = services.calculate_attendance_summary(["Present"] * 4)
        self.assertEqual(summary["percentage"], 100.0)

    def test_late_counts_as_present_but_absent_and_excused_do_not(self):
        summary = services.calculate_attendance_summary(
            ["Present", "Late", "Absent", "Excused"]
        )
        self.assertEqual(
            summary, {"totalDays": 4, "presentDays": 2, "percentage": 50.0}
        )

    def test_percentage_is_rounded_to_one_decimal(self):
        summary = services.calculate_attendance_summary(["Present", "Present", "Absent"])
        self.assertEqual(summary["percentage"], 66.7)


class AttendanceSummaryApiTests(AttendanceTestBase):
    """Admin attendance totals grouped by section and academic class."""

    url = "/api/attendance/summary"

    def test_selected_period_returns_class_and_grade_totals(self):
        extra_section = Section.objects.create(academic_class=self.cls, name="B")
        extra_student = Student.objects.create(
            name="Student 4",
            student_class="Grade 10",
            roll_number=4,
            contact_details="parent4@example.com",
            section=extra_section,
        )
        other_grade = AcademicClass.objects.create(name="Grade 9")
        Section.objects.create(academic_class=other_grade, name="A")

        services.mark_attendance(self.teacher, self.cls.pk, "2026-10-19", [])
        services.mark_attendance(self.teacher, self.cls.pk, "2026-10-20", [])
        services.mark_attendance(
            self.teacher,
            self.cls.pk,
            "2026-10-21",
            [
                {"studentId": self.students[0].pk, "status": "Absent"},
                {"studentId": self.students[1].pk, "status": "Late"},
                {"studentId": self.students[2].pk, "status": "Excused"},
            ],
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            self.url, {"startDate": "2026-10-20", "endDate": "2026-10-21"}
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(
            (body["startDate"], body["endDate"]),
            ("2026-10-20", "2026-10-21"),
        )
        classes = {item["className"]: item for item in body["classes"]}
        self.assertEqual(
            classes["Grade 10 - A"],
            {
                "className": "Grade 10 - A",
                "totalRecords": 6,
                "present": 3,
                "absent": 1,
                "late": 1,
                "excused": 1,
                "averagePercentage": 66.7,
            },
        )
        self.assertEqual(classes["Grade 10 - B"]["totalRecords"], 2)
        self.assertEqual(classes["Grade 10 - B"]["averagePercentage"], 100.0)
        grades = {item["gradeName"]: item for item in body["grades"]}
        self.assertEqual(
            grades["Grade 10"],
            {
                "gradeName": "Grade 10",
                "totalRecords": 8,
                "present": 5,
                "absent": 1,
                "late": 1,
                "excused": 1,
                "averagePercentage": 75.0,
            },
        )
        self.assertEqual(grades["Grade 9"]["totalRecords"], 0)
        self.assertEqual(grades["Grade 9"]["averagePercentage"], 0.0)

    def test_period_without_attendance_returns_zero_totals(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            self.url, {"startDate": "2026-10-24", "endDate": "2026-10-25"}
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertTrue(body["classes"])
        self.assertTrue(body["grades"])
        for item in body["classes"] + body["grades"]:
            self.assertEqual(item["totalRecords"], 0)
            self.assertEqual(item["averagePercentage"], 0.0)

    def test_invalid_or_reversed_date_range_is_rejected(self):
        self.client.force_login(self.admin)
        for params in (
            {"startDate": "not-a-date", "endDate": "2026-10-25"},
            {"startDate": "2026-10-26", "endDate": "2026-10-25"},
        ):
            with self.subTest(params=params):
                response = self.client.get(self.url, params)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.json()["status"], "error")

    def test_summary_is_admin_only(self):
        for user in (self.teacher,):
            with self.subTest(role=user.role):
                self.client.force_login(user)
                self.assertEqual(
                    self.client.get(
                        self.url,
                        {"startDate": "2026-10-20", "endDate": "2026-10-21"},
                    ).status_code,
                    403,
                )

    def test_summary_filters_by_date_class_and_student_with_and_semantics(self):
        other_class = AcademicClass.objects.create(name="Grade 9")
        other_section = Section.objects.create(academic_class=other_class, name="A")
        other_student = Student.objects.create(
            name="Other student",
            student_class="Grade 9",
            roll_number=90,
            contact_details="other@example.com",
            section=other_section,
        )
        services.mark_attendance(self.teacher, self.cls.pk, "2026-10-19", [])
        services.mark_attendance(
            self.teacher,
            self.cls.pk,
            "2026-10-20",
            [{"studentId": self.students[0].pk, "status": "Absent"}],
        )
        Attendance.objects.create(
            student=other_student,
            academic_class=other_class,
            date="2026-10-20",
            status="Present",
        )

        self.client.force_login(self.admin)
        params = {
            "startDate": "2026-10-20",
            "endDate": "2026-10-20",
            "classId": self.cls.pk,
            "studentId": self.students[0].pk,
        }
        response = self.client.get(self.url, params)
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["classes"][0]["totalRecords"], 1)
        self.assertEqual(body["classes"][0]["absent"], 1)
        self.assertEqual(body["classes"][0]["averagePercentage"], 0.0)
        self.assertEqual(body["grades"][0]["totalRecords"], 1)

    def test_summary_class_and_student_filters_can_be_used_individually(self):
        services.mark_attendance(
            self.teacher,
            self.cls.pk,
            "2026-10-20",
            [{"studentId": self.students[0].pk, "status": "Absent"}],
        )
        self.client.force_login(self.admin)
        class_response = self.client.get(
            self.url,
            {"startDate": "2026-10-20", "endDate": "2026-10-20", "classId": self.cls.pk},
        )
        self.assertEqual(class_response.status_code, 200)
        self.assertEqual(class_response.json()["grades"][0]["totalRecords"], 3)
        student_response = self.client.get(
            self.url,
            {"startDate": "2026-10-20", "endDate": "2026-10-20", "studentId": self.students[0].pk},
        )
        self.assertEqual(student_response.status_code, 200)
        self.assertEqual(student_response.json()["grades"][0]["totalRecords"], 1)

    def test_summary_filters_return_empty_scoped_totals(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            self.url,
            {
                "startDate": "2026-10-20",
                "endDate": "2026-10-20",
                "classId": self.cls.pk,
                "studentId": self.students[0].pk,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["classes"][0]["totalRecords"], 0)
        self.assertEqual(response.json()["grades"][0]["averagePercentage"], 0.0)

    def test_summary_rejects_invalid_filter_identifiers(self):
        self.client.force_login(self.admin)
        for key, value in (("classId", "not-a-number"), ("studentId", "not-a-number")):
            response = self.client.get(
                self.url,
                {"startDate": TODAY, "endDate": TODAY, key: value},
            )
            with self.subTest(key=key):
                self.assertEqual(response.status_code, 400)
        for key in ("classId", "studentId"):
            response = self.client.get(
                self.url,
                {"startDate": TODAY, "endDate": TODAY, key: 9999},
            )
            with self.subTest(missing=key):
                self.assertEqual(response.status_code, 404)


class AttendanceReportExportTests(AttendanceTestBase):
    base_url = "/api/attendance/reports"

    def export(self, report, file_format, **params):
        return self.client.get(
            f"{self.base_url}/{report}/{file_format}",
            params,
        )

    def test_summary_excel_contains_displayed_class_and_grade_rows(self):
        services.mark_attendance(
            self.teacher,
            self.cls.pk,
            TODAY,
            [{"studentId": self.students[1].pk, "status": "Absent"}],
        )
        self.client.force_login(self.admin)
        response = self.export(
            "summary",
            "excel",
            startDate=TODAY,
            endDate=TODAY,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        class_sheet = workbook["By class"]
        rows = list(class_sheet.iter_rows(values_only=True))
        headers = rows[3]
        self.assertEqual(
            rows[4],
            ("Grade 10 - A", 3, 2, 1, 0, 0, "66.7%"),
        )
        self.assertEqual(headers[0], "Class")
        grade_rows = list(workbook["By grade"].iter_rows(values_only=True))
        self.assertEqual(grade_rows[4][0:2], ("Grade 10", 3))

    def test_daily_excel_contains_the_records_shown_in_daily_view(self):
        services.mark_attendance(
            self.teacher,
            self.cls.pk,
            TODAY,
            [{"studentId": self.students[1].pk, "status": "Absent"}],
        )
        self.client.force_login(self.teacher)
        response = self.export(
            "daily",
            "excel",
            classId=self.cls.pk,
            date=TODAY,
        )

        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(BytesIO(response.content), data_only=True)["Attendance records"]
        rows = list(worksheet.iter_rows(values_only=True))
        header_index = rows.index(("Roll no.", "Name", "Status", "Remarks"))
        self.assertEqual(
            [row[:3] for row in rows[header_index + 1:header_index + 4]],
            [(1, "Student 1", "Present"), (2, "Student 2", "Absent"), (3, "Student 3", "Present")],
        )

    def test_student_excel_exports_empty_history_with_headers(self):
        self.client.force_login(self.admin)
        response = self.export("student", "excel", studentId=self.students[0].pk)

        self.assertEqual(response.status_code, 200)
        rows = list(
            load_workbook(BytesIO(response.content), data_only=True)
            ["Attendance history"]
            .iter_rows(values_only=True)
        )
        header_index = rows.index(("Date", "Status", "Remarks"))
        self.assertEqual(rows[header_index + 1][0], "No records for this report.")

    def test_pdf_exports_are_valid_and_empty_daily_report_is_downloadable(self):
        self.client.force_login(self.teacher)
        response = self.export("daily", "pdf", classId=self.cls.pk, date=TODAY)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn('filename="attendance-daily.pdf"', response["Content-Disposition"])

    def test_summary_and_student_reports_can_be_exported_as_pdf(self):
        parent = User.objects.create_user(
            username="export-parent",
            password="test-password",
            role=User.Role.STUDENT_PARENT,
        )
        self.students[0].parent = parent
        self.students[0].save()
        services.mark_attendance(self.teacher, self.cls.pk, TODAY, [])

        self.client.force_login(self.admin)
        summary = self.export(
            "summary", "pdf", startDate=TODAY, endDate=TODAY
        )
        self.assertEqual(summary.status_code, 200)
        self.assertEqual(summary["Content-Type"], "application/pdf")
        self.assertTrue(summary.content.startswith(b"%PDF-"))

        self.client.force_login(parent)
        history = self.export("student", "pdf", studentId=self.students[0].pk)
        self.assertEqual(history.status_code, 200)
        self.assertEqual(history["Content-Type"], "application/pdf")
        self.assertTrue(history.content.startswith(b"%PDF-"))

    def test_role_restrictions_are_preserved_for_exports(self):
        self.client.force_login(self.teacher)
        summary = self.export(
            "summary", "excel", startDate=TODAY, endDate=TODAY
        )
        self.assertEqual(summary.status_code, 403)

        parent = User.objects.create_user(
            username="export-parent",
            password="test-password",
            role=User.Role.STUDENT_PARENT,
        )
        self.client.force_login(parent)
        daily = self.export("daily", "pdf", classId=self.cls.pk, date=TODAY)
        self.assertEqual(daily.status_code, 403)

    def test_unsupported_format_and_unknown_report_are_reported(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.export("summary", "csv").status_code, 400)
        self.assertEqual(self.export("unknown", "pdf").status_code, 404)


class StudentHistoryApiTests(AttendanceTestBase):
    """REQ4.2: a student's history and percentage, visible only to the right people."""

    def setUp(self):
        super().setUp()
        self.parent = User.objects.create_user(
            username="parent1", password="pass12345", role=User.Role.STUDENT_PARENT
        )
        self.other_parent = User.objects.create_user(
            username="parent2", password="pass12345", role=User.Role.STUDENT_PARENT
        )
        self.child, self.other_child, self.unlinked_child = self.students
        self.child.parent = self.parent
        self.child.save()
        self.other_child.parent = self.other_parent
        self.other_child.save()

        # Day 1: everyone present. Day 2: child Absent. Day 3: child Late.
        services.mark_attendance(self.teacher, self.cls.pk, "2026-10-20", [])
        services.mark_attendance(
            self.teacher, self.cls.pk, "2026-10-21",
            [{"studentId": self.child.pk, "status": "Absent"}],
        )
        services.mark_attendance(
            self.teacher, self.cls.pk, "2026-10-22",
            [{"studentId": self.child.pk, "status": "Late", "remarks": "Bus"}],
        )

    def history(self, student):
        return self.client.get(f"/api/attendance/student/{student.pk}")

    def test_parent_sees_their_own_childs_history_and_percentage(self):
        self.client.force_login(self.parent)
        response = self.history(self.child)
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["student"]["name"], "Student 1")
        self.assertEqual(body["totalDays"], 3)
        self.assertEqual(body["presentDays"], 2)
        self.assertEqual(body["percentage"], 66.7)
        self.assertEqual(
            body["counts"], {"Present": 1, "Absent": 1, "Late": 1, "Excused": 0}
        )
        # Newest first
        self.assertEqual(
            [(r["date"], r["status"]) for r in body["records"]],
            [("2026-10-22", "Late"), ("2026-10-21", "Absent"), ("2026-10-20", "Present")],
        )
        self.assertEqual(body["records"][0]["remarks"], "Bus")

    def test_parent_cannot_see_another_students_data(self):
        self.client.force_login(self.parent)
        self.assertEqual(self.history(self.other_child).status_code, 403)

    def test_parent_cannot_see_a_student_who_is_not_linked_to_anyone(self):
        self.client.force_login(self.parent)
        self.assertEqual(self.history(self.unlinked_child).status_code, 403)

    def test_parent_with_no_linked_children_gets_403(self):
        stranger = User.objects.create_user(
            username="parent3", password="pass12345", role=User.Role.STUDENT_PARENT
        )
        self.client.force_login(stranger)
        self.assertEqual(self.history(self.child).status_code, 403)

    def test_admin_can_see_any_student(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.history(self.unlinked_child).status_code, 200)

    def test_class_teacher_can_see_their_student_but_other_teacher_cannot(self):
        self.client.force_login(self.teacher)
        self.assertEqual(self.history(self.child).status_code, 200)
        self.client.force_login(self.other_teacher)
        self.assertEqual(self.history(self.child).status_code, 403)

    def test_not_logged_in_is_refused(self):
        self.assertIn(self.history(self.child).status_code, (401, 403))

    def test_unknown_student_gets_404(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get("/api/attendance/student/9999").status_code, 404)

    def test_student_with_no_records_returns_zero_days(self):
        newcomer = Student.objects.create(
            name="Newcomer", student_class="Grade 10", roll_number=50,
            contact_details="new@example.com", section=self.child.section,
            parent=self.parent,
        )
        self.client.force_login(self.parent)
        body = self.history(newcomer).json()
        self.assertEqual((body["totalDays"], body["percentage"], body["records"]), (0, 0.0, []))

    def test_history_date_and_class_filters_recalculate_percentage(self):
        self.client.force_login(self.parent)
        response = self.client.get(
            f"/api/attendance/student/{self.child.pk}",
            {
                "startDate": "2026-10-21",
                "endDate": "2026-10-22",
                "classId": self.cls.pk,
            },
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["totalDays"], 2)
        self.assertEqual(body["presentDays"], 1)
        self.assertEqual(body["percentage"], 50.0)
        self.assertEqual(len(body["records"]), 2)

    def test_history_filters_validate_dates_and_class(self):
        self.client.force_login(self.parent)
        response = self.client.get(
            f"/api/attendance/student/{self.child.pk}",
            {"startDate": "bad-date", "endDate": "2026-10-22"},
        )
        self.assertEqual(response.status_code, 400)
        response = self.client.get(
            f"/api/attendance/student/{self.child.pk}",
            {"startDate": "2026-10-21", "endDate": "2026-10-20"},
        )
        self.assertEqual(response.status_code, 400)
        response = self.client.get(
            f"/api/attendance/student/{self.child.pk}",
            {"classId": "bad-class"},
        )
        self.assertEqual(response.status_code, 400)
        response = self.client.get(
            f"/api/attendance/student/{self.child.pk}",
            {"classId": 9999},
        )
        self.assertEqual(response.status_code, 404)