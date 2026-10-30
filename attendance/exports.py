"""Generate downloadable Excel and PDF attendance reports."""

from io import BytesIO
from pathlib import Path

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font


def _safe_sheet_name(name):
    for character in "[]:*?/\\":
        name = name.replace(character, " ")
    return name[:31] or "Report"


def _append_excel_row(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _excel_response(title, metadata, tables, filename):
    workbook = Workbook()
    workbook.remove(workbook.active)
    bold = Font(bold=True)

    for table_title, headers, rows in tables:
        sheet = workbook.create_sheet(_safe_sheet_name(table_title))
        _append_excel_row(sheet, [title])
        sheet["A1"].font = Font(bold=True, size=14)
        for label, value in metadata:
            _append_excel_row(sheet, [label, value])
        _append_excel_row(sheet, [""])
        header_row = sheet.max_row + 1
        _append_excel_row(sheet, headers)
        for cell in sheet[header_row]:
            cell.font = bold
        for row in rows:
            _append_excel_row(sheet, list(row))
        if not rows:
            _append_excel_row(sheet, ["No records for this report."])
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = (
            f"A{header_row}:{sheet.cell(sheet.max_row, len(headers)).coordinate}"
        )
        for column in sheet.columns:
            width = max(
                (len(str(cell.value or "")) for cell in column),
                default=0,
            )
            sheet.column_dimensions[column[0].column_letter].width = min(width + 2, 40)

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def _register_pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    from reportlab.pdfbase.ttfonts import TTFont

    if "SAMS-Vera" in pdfmetrics.getRegisteredFontNames():
        return
    from reportlab import __file__ as reportlab_path

    font_dir = Path(reportlab_path).parent / "fonts"
    pdfmetrics.registerFont(TTFont("SAMS-Vera", str(font_dir / "Vera.ttf")))
    pdfmetrics.registerFont(TTFont("SAMS-Vera-Bold", str(font_dir / "VeraBd.ttf")))
    registerFontFamily(
        "SAMS-Vera",
        normal="SAMS-Vera",
        bold="SAMS-Vera-Bold",
        italic="SAMS-Vera",
        boldItalic="SAMS-Vera-Bold",
    )


def _pdf_response(title, metadata, tables, filename):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    _register_pdf_fonts()
    output = BytesIO()
    page_size = landscape(A4) if any(len(headers) > 5 for _, headers, _ in tables) else A4
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    styles["Title"].fontName = "SAMS-Vera-Bold"
    styles["Normal"].fontName = "SAMS-Vera"
    styles["Heading2"].fontName = "SAMS-Vera-Bold"
    header_style = ParagraphStyle(
        "ReportHeader",
        parent=styles["Normal"],
        fontName="SAMS-Vera-Bold",
        textColor=colors.white,
    )

    story = [Paragraph(escape(title), styles["Title"])]
    story.extend(
        Paragraph(f"<b>{escape(str(label))}:</b> {escape(str(value))}", styles["Normal"])
        for label, value in metadata
    )
    story.append(Spacer(1, 6 * mm))

    for table_title, headers, rows in tables:
        story.append(Paragraph(escape(table_title), styles["Heading2"]))
        values = [[Paragraph(escape(str(value)), header_style) for value in headers]]
        values.extend(
            [Paragraph(escape(str(value)), styles["Normal"]) for value in row]
            for row in rows
        )
        column_width = document.width / len(headers)
        report_table = Table(values, repeatRows=1, colWidths=[column_width] * len(headers))
        report_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#212529")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "SAMS-Vera-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "SAMS-Vera"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.whitesmoke]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(report_table)
        if not rows:
            story.append(Paragraph("No records for this report.", styles["Normal"]))
        story.append(Spacer(1, 5 * mm))

    document.build(story)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return response


def build_report_response(file_format, title, metadata, tables, filename):
    """Return a workbook or PDF containing the same table data as a report."""
    if file_format == "excel":
        return _excel_response(title, metadata, tables, filename)
    if file_format == "pdf":
        return _pdf_response(title, metadata, tables, filename)
    raise ValueError(f"Unsupported report format: {file_format}")
