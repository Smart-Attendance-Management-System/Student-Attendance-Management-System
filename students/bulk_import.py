import csv
import io
import zipfile
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .forms import StudentForm
from .models import AcademicClass, Section, Student

MAX_UPLOAD_SIZE = 5 * 1024 * 1024
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}
REQUIRED_COLUMNS = {"name", "class", "roll_number", "contact_details"}
OPTIONAL_COLUMNS = {"section"}
COLUMN_ALIASES = {
    "student_class": "class",
    "class_name": "class",
    "roll": "roll_number",
    "roll_no": "roll_number",
    "contact": "contact_details",
    "contact_detail": "contact_details",
}


class BulkImportError(Exception):
    """A safe, user-facing error for a rejected upload."""


def _normalise_header(value):
    return " ".join(str(value or "").strip().lower().replace("_", " ").split()).replace(
        " ", "_"
    )


def _normalise_value(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_csv(uploaded_file):
    try:
        content = uploaded_file.read().decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BulkImportError("The CSV file must use UTF-8 encoding.") from exc
    return list(csv.reader(io.StringIO(content)))


def _read_xlsx(uploaded_file):
    workbook = None
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        worksheet = workbook.active
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise BulkImportError("The Excel file could not be read.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _read_rows(uploaded_file):
    suffix = Path(uploaded_file.name or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise BulkImportError("Upload a CSV file or an .xlsx Excel file.")
    if uploaded_file.size > MAX_UPLOAD_SIZE:
        raise BulkImportError("The uploaded file must be 5 MB or smaller.")
    if uploaded_file.size == 0:
        raise BulkImportError("The uploaded file is empty.")
    if suffix == ".csv":
        return _read_csv(uploaded_file)
    return _read_xlsx(uploaded_file)


def _map_rows(raw_rows):
    if not raw_rows:
        raise BulkImportError("The uploaded file does not contain any rows.")

    headers = [_normalise_header(value) for value in raw_rows[0]]
    headers = [COLUMN_ALIASES.get(header, header) for header in headers]
    if not any(headers):
        raise BulkImportError("The uploaded file is missing its header row.")
    if len(headers) != len(set(headers)):
        raise BulkImportError("The uploaded file contains duplicate column headers.")

    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        missing_labels = ", ".join(sorted(missing))
        raise BulkImportError(f"Required columns are missing: {missing_labels}.")
    unexpected = set(headers) - REQUIRED_COLUMNS - OPTIONAL_COLUMNS - {""}
    if unexpected:
        unexpected_labels = ", ".join(sorted(unexpected))
        raise BulkImportError(f"Unsupported columns were found: {unexpected_labels}.")

    rows = []
    for row_number, values in enumerate(raw_rows[1:], start=2):
        data = {
            header: _normalise_value(values[index] if index < len(values) else "")
            for index, header in enumerate(headers)
            if header
        }
        if not any(data.values()):
            continue
        data["row_number"] = row_number
        rows.append(data)
    if not rows:
        raise BulkImportError("The uploaded file does not contain any student records.")
    return rows


def _resolve_section(data, row_number):
    class_name = data["class"]
    academic_class = AcademicClass.objects.filter(name__iexact=class_name).first()
    if academic_class is None:
        raise BulkImportError(f"Row {row_number}: class '{class_name}' was not found.")

    section_name = data.get("section", "")
    if not section_name:
        return academic_class, None
    section = Section.objects.filter(
        academic_class=academic_class,
        name__iexact=section_name,
    ).first()
    if section is None:
        raise BulkImportError(
            f"Row {row_number}: section '{section_name}' was not found in class '{class_name}'."
        )
    return academic_class, section


def import_students(uploaded_file):
    rows = _map_rows(_read_rows(uploaded_file))
    seen_roll_numbers = set()
    forms = []
    failures = []

    for data in rows:
        row_number = data.pop("row_number")
        for field, label in (
            ("name", "name"),
            ("class", "class"),
            ("roll_number", "roll number"),
            ("contact_details", "contact details"),
        ):
            if not data.get(field):
                failures.append(
                    {"row_number": row_number, "reason": f"{label} is required."}
                )
                break
        else:
            try:
                roll_number = int(data["roll_number"])
            except (TypeError, ValueError):
                failures.append(
                    {
                        "row_number": row_number,
                        "reason": "roll number must be a positive whole number.",
                    }
                )
                continue
            if roll_number <= 0:
                failures.append(
                    {
                        "row_number": row_number,
                        "reason": "roll number must be a positive whole number.",
                    }
                )
                continue
            if roll_number in seen_roll_numbers:
                failures.append(
                    {
                        "row_number": row_number,
                        "reason": f"roll number {roll_number} is duplicated in the file.",
                    }
                )
                continue
            seen_roll_numbers.add(roll_number)
            if Student.objects.filter(roll_number=roll_number).exists():
                failures.append(
                    {
                        "row_number": row_number,
                        "reason": f"roll number {roll_number} already exists.",
                    }
                )
                continue

            try:
                academic_class, section = _resolve_section(data, row_number)
            except BulkImportError as exc:
                failures.append(
                    {"row_number": row_number, "reason": str(exc).split(": ", 1)[-1]}
                )
                continue

            form = StudentForm(
                {
                    "name": data.get("name", ""),
                    "student_class": str(academic_class),
                    "section": section.pk if section else "",
                    "roll_number": roll_number,
                    "contact_details": data.get("contact_details", ""),
                }
            )
            if not form.is_valid():
                errors = "; ".join(
                    f"{field}: {', '.join(messages)}"
                    for field, messages in form.errors.items()
                )
                failures.append({"row_number": row_number, "reason": errors})
                continue
            forms.append(form)

    try:
        with transaction.atomic():
            students = [form.save(commit=False) for form in forms]
            for student in students:
                student.full_clean()
            Student.objects.bulk_create(students)
    except (ValidationError, IntegrityError) as exc:
        raise BulkImportError(
            "The import could not be completed because one or more records are invalid."
        ) from exc
    return {
        "imported_count": len(forms),
        "failed_count": len(failures),
        "failures": failures,
    }
